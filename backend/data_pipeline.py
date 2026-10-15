from __future__ import annotations

import hashlib
import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    from .database import TRANSACTION_COLUMNS, init_db, replace_transactions
except ImportError:  # Supports running scripts from inside backend/ as well.
    from database import TRANSACTION_COLUMNS, init_db, replace_transactions

DEFAULT_WORKBOOK = Path(__file__).resolve().with_name("Database_Core_Pendapatan_TTH_2020_2025.xlsx")
AMOUNT_COLUMNS = {"nominal_revenue", "ppn", "pph23", "nominal_lost"}
MONTHS = ("Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember")


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def normalize_cell(column: str, value: Any) -> Any:
    if value is None:
        return 0.0 if column in AMOUNT_COLUMNS else ""
    if column == "tahun":
        try:
            year = int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Invalid year value in workbook: {value!r}") from exc
        if year < 1900 or year > 2200:
            raise ValueError(f"Year value outside allowed range: {year}")
        return year
    if column in AMOUNT_COLUMNS:
        try:
            amount = float(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Invalid amount in column '{column}': {value!r}") from exc
        if not math.isfinite(amount):
            raise ValueError(f"Non-finite amount in column '{column}'")
        return amount
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def derive_transaction_fields(
    jenis_transaksi: str,
    pengujian_stel: str,
    laboratorium: str,
    keterangan: str,
    pph23: float,
) -> dict[str, Any]:
    jenis_upper = jenis_transaksi.upper()
    pengujian_upper = pengujian_stel.upper()
    lab_upper = laboratorium.upper()
    notes_lower = keterangan.lower()

    kategori_utama = "STEL" if "STEL" in jenis_upper or "STEL" in pengujian_upper else "Pengujian"
    if "KAL" in lab_upper or "KAL" in pengujian_upper or "KALIBRASI" in lab_upper:
        unit_org = "Kalibrasi"
    elif any(token in lab_upper for token in ("DES", "DEV", "SSA", "ENG")):
        unit_org = "DES"
    else:
        unit_org = "IQA"

    if "QA" in pengujian_upper:
        kategori_uji = "QA"
    elif "TA" in pengujian_upper:
        kategori_uji = "TA"
    elif "VT" in pengujian_upper:
        kategori_uji = "VT"
    else:
        kategori_uji = "Lainnya"

    if "DEV" in lab_upper or "DEVICE" in lab_upper:
        domain_lab = "Device"
    elif any(token in lab_upper for token in ("ENG", "ENERGY", "SSA")):
        domain_lab = "Energy"
    elif any(token in lab_upper for token in ("KABEL", "FBA", "OPTIC")):
        domain_lab = "FBA/Kabel"
    elif any(token in lab_upper for token in ("TRANSMISI", "TRS", "RADIO")):
        domain_lab = "Transmisi"
    else:
        domain_lab = "Lainnya/General"

    unpaid_terms = ("tidak membaya", "tidak bayar", "tidak banyar", "unpaid")
    if any(term in notes_lower for term in unpaid_terms):
        status_lost = "Unpaid PPh 23"
        nominal_lost = abs(pph23)
    elif any(term in notes_lower for term in ("kurang", "belum")):
        status_lost = "Kurang/Belum Bayar"
        nominal_lost = abs(pph23)
    else:
        status_lost = "No Loss Identified"
        nominal_lost = 0.0

    return {
        "kategori_utama": kategori_utama,
        "unit_org": unit_org,
        "kategori_uji": kategori_uji,
        "domain_lab": domain_lab,
        "status_lost": status_lost,
        "nominal_lost": nominal_lost,
    }


def rows_from_workbook(workbook_path: Path = DEFAULT_WORKBOOK) -> list[tuple[Any, ...]]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Master Data" not in workbook.sheetnames:
            raise ValueError("Workbook must include the 'Master Data' sheet.")
        worksheet = workbook["Master Data"]
        iterator = worksheet.iter_rows(values_only=True)
        header_cells = next(iterator, None)
        if not header_cells:
            raise ValueError("The 'Master Data' sheet is empty.")

        header_map = {normalize_header(name): index for index, name in enumerate(header_cells) if name is not None}
        missing_columns = [name for name in TRANSACTION_COLUMNS if name not in header_map]
        if missing_columns:
            raise ValueError(f"Workbook is missing required columns: {', '.join(missing_columns)}")

        records: list[tuple[Any, ...]] = []
        for row_number, row in enumerate(iterator, start=2):
            if not row or not any(value is not None for value in row):
                continue
            record = tuple(
                normalize_cell(column, row[header_map[column]] if header_map[column] < len(row) else None)
                for column in TRANSACTION_COLUMNS
            )
            if not record[0] or not record[1] or not record[3]:
                raise ValueError(f"Required year, month, or company is missing at workbook row {row_number}.")
            records.append(record)

        if not records:
            raise ValueError("No transaction rows were found; refusing to replace the live database.")
        return records
    finally:
        workbook.close()


def _pick_column(headers: dict[str, int], *needles: str) -> int | None:
    for needle in needles:
        for name, index in headers.items():
            if needle in name:
                return index
    return None


def _value(row: tuple[Any, ...], index: int | None, default: Any = "") -> Any:
    if index is None or index >= len(row) or row[index] is None:
        return default
    return row[index]


def _infer_year(raw_date: Any, workbook_path: Path, year_override: int | None) -> int:
    if year_override is not None:
        return year_override
    if isinstance(raw_date, (datetime, date)):
        return raw_date.year
    date_year = re.search(r"(?<!\d)(20\d{2})(?!\d)", str(raw_date))
    filename_year = re.search(r"(?<!\d)(20\d{2})(?!\d)", workbook_path.stem)
    match = date_year or filename_year
    if match:
        return int(match.group(1))
    raise ValueError(
        f"Could not infer the transaction year in '{workbook_path.name}'. "
        "Pass the year query parameter or include a year in the filename."
    )


def _amount(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        amount = float(value)
        return amount if math.isfinite(amount) else 0.0

    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"[^0-9,.+-]", "", text.strip("() "))
    if not text:
        return 0.0

    dot, comma = text.rfind("."), text.rfind(",")
    if dot >= 0 and comma >= 0:
        decimal_separator = "." if dot > comma else ","
        thousands_separator = "," if decimal_separator == "." else "."
        text = text.replace(thousands_separator, "").replace(decimal_separator, ".")
    elif comma >= 0:
        decimals = len(text) - comma - 1
        text = text.replace(",", "." if decimals in (1, 2) else "")
    elif dot >= 0:
        decimals = len(text) - dot - 1
        if decimals == 3 and dot > 0:
            text = text.replace(".", "")

    try:
        amount = float(text)
    except ValueError:
        return 0.0
    if negative:
        amount = -abs(amount)
    return amount if math.isfinite(amount) else 0.0


def rows_from_monthly_workbook(
    workbook_path: Path,
    year_override: int | None = None,
    source_name: str | None = None,
) -> list[tuple[Any, ...]]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    records: list[tuple[Any, ...]] = []
    try:
        for month in MONTHS:
            if month not in workbook.sheetnames:
                continue
            worksheet = workbook[month]
            iterator = worksheet.iter_rows(values_only=True)
            header_row = next(iterator, None)
            if not header_row:
                continue
            headers = {normalize_header(value): index for index, value in enumerate(header_row) if value is not None}
            company_idx = _pick_column(headers, "perusahaan")
            if company_idx is None:
                continue

            date_idx = _pick_column(headers, "tanggal", "unnamed")
            revenue_idx = _pick_column(headers, "nilai_transfer", "nominal_revenue", "jumlah", "transfer")
            pph_idx = _pick_column(headers, "pph")
            ppn_idx = _pick_column(headers, "ppn")
            type_idx = _pick_column(headers, "jenis")
            test_idx = _pick_column(headers, "pengujian")
            lab_idx = _pick_column(headers, "laboratorium")
            notes_idx = _pick_column(headers, "keterangan", "ket")

            for row in iterator:
                company = str(_value(row, company_idx)).strip()
                if not company or re.fullmatch(r"(?:total|jumlah|rekap)(?:\s+.*)?", company, re.IGNORECASE):
                    continue
                revenue = _amount(_value(row, revenue_idx, 0))
                pph23 = _amount(_value(row, pph_idx, 0))
                ppn = _amount(_value(row, ppn_idx, 0))
                transaction_type = str(_value(row, type_idx, "Pengujian Perangkat")).strip()
                test_type = str(_value(row, test_idx, "QA")).strip()
                lab = str(_value(row, lab_idx, "Unspecified")).strip()
                notes = str(_value(row, notes_idx, "")).strip()
                raw_date = _value(row, date_idx, "")
                date_value = normalize_cell("tanggal", raw_date)
                year = _infer_year(raw_date, Path(source_name) if source_name else workbook_path, year_override)
                derived = derive_transaction_fields(transaction_type, test_type, lab, notes, pph23)
                row_values = {
                    "tahun": year,
                    "bulan": month,
                    "tanggal": date_value,
                    "perusahaan": company,
                    "jenis_transaksi": transaction_type,
                    "pengujian_stel": test_type,
                    "laboratorium": lab,
                    "keterangan": notes,
                    "nominal_revenue": revenue,
                    "ppn": ppn,
                    "pph23": pph23,
                    **derived,
                }
                records.append(tuple(row_values[column] for column in TRANSACTION_COLUMNS))
    finally:
        workbook.close()

    return records


def file_sha256(workbook_path: Path) -> str:
    digest = hashlib.sha256()
    with workbook_path.open("rb") as workbook_file:
        for chunk in iter(lambda: workbook_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def import_workbook(workbook_path: Path = DEFAULT_WORKBOOK) -> tuple[int, Path | None]:
    records = rows_from_workbook(workbook_path)
    init_db()
    imported_count, backup_path = replace_transactions(
        records,
        source_hash=file_sha256(workbook_path),
        source_name=workbook_path.name,
        make_backup=True,
    )
    return imported_count, backup_path
