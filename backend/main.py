from __future__ import annotations

import hashlib
import logging
import os
import re
import tempfile
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, RedirectResponse
from pydantic import BaseModel, ConfigDict, Field

try:
    from .database import PROJECT_ROOT, TRANSACTION_COLUMNS, DuplicateWorkbookError, database_session, init_db, insert_transactions, replace_transactions
    from .data_pipeline import (
        derive_transaction_fields,
        file_sha256,
        rows_from_monthly_workbook,
        rows_from_workbook,
    )
except ImportError:  # Supports `python main.py` from inside backend/.
    from database import PROJECT_ROOT, TRANSACTION_COLUMNS, DuplicateWorkbookError, database_session, init_db, insert_transactions, replace_transactions
    from data_pipeline import (
        derive_transaction_fields,
        file_sha256,
        rows_from_monthly_workbook,
        rows_from_workbook,
    )

app = FastAPI(
    title="Telkom Revenue Database API",
    description="Database API for the Telkom revenue workspace.",
    version="2.0.0",
)

FRONTEND_ORIGINS = [
    origin.strip()
    for origin in os.getenv(
        "FRONTEND_ORIGINS",
        "http://localhost:3000,http://127.0.0.1:3000,https://telkom-intern-project.vercel.app",
    ).split(",")
    if origin.strip()
]
logger = logging.getLogger(__name__)

app.add_middleware(
    CORSMiddleware,
    allow_origins=FRONTEND_ORIGINS,
    allow_origin_regex=r"https://.*\.vercel\.app",
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

MAX_UPLOAD_BYTES = 50 * 1024 * 1024
UPLOAD_DIR = PROJECT_ROOT / "uploads"


class TransactionInput(BaseModel):
    model_config = ConfigDict(extra="forbid")

    tahun: int = Field(ge=1900, le=2200)
    bulan: str = Field(min_length=1, max_length=32)
    tanggal: str = ""
    perusahaan: str = Field(min_length=1, max_length=300)
    jenis_transaksi: str = Field(default="", max_length=200)
    pengujian_stel: str = Field(default="", max_length=200)
    laboratorium: str = Field(default="", max_length=200)
    keterangan: str = ""
    nominal_revenue: float = Field(ge=0)
    ppn: float = Field(default=0, ge=0)
    pph23: float = 0


class ChatQuery(BaseModel):
    prompt: str = Field(min_length=1, max_length=2000)


CORE_WORKBOOK = Path(__file__).resolve().parent / "Database_Core_Pendapatan_TTH_2020_2025.xlsx"


@app.on_event("startup")
def startup_database() -> None:
    init_db()
    # Auto-seed dari workbook inti saat deploy pertama (database kosong)
    with database_session() as connection:
        count = connection.execute("SELECT COUNT(*) FROM transactions").fetchone()[0]
    if count == 0 and CORE_WORKBOOK.exists():
        logger.info("Database kosong — memuat data dari workbook inti: %s", CORE_WORKBOOK.name)
        try:
            imported = process_excel_file(CORE_WORKBOOK)
            logger.info("Berhasil memuat %d transaksi dari workbook inti.", imported)
        except Exception as exc:
            logger.error("Gagal memuat workbook inti saat startup: %s", exc)


def file_sha256_for_year(source_hash: str, year: int) -> str:
    return hashlib.sha256(f"{source_hash}:year:{year}".encode("utf-8")).hexdigest()


def process_excel_file(
    file_path: str | Path,
    year_override: int | None = None,
    source_name: str | None = None,
) -> int:
    """Read either the normalized core workbook or the monthly upload format."""
    from openpyxl import load_workbook

    path = Path(file_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        has_master_data = "Master Data" in workbook.sheetnames
    finally:
        workbook.close()

    source_hash = file_sha256(path)
    if has_master_data:
        rows = rows_from_workbook(path)
        if year_override is None:
            imported_count, _ = replace_transactions(
                rows,
                source_hash=source_hash,
                source_name=path.name,
                make_backup=True,
            )
            return imported_count
        rows = [row for row in rows if row[0] == year_override]
        source_hash = file_sha256_for_year(source_hash, year_override)
    else:
        rows = rows_from_monthly_workbook(
            path,
            year_override,
            source_name=source_name or path.name,
        )
    return insert_transactions(
        rows,
        source_hash=source_hash,
        source_name=source_name or path.name,
    )


def _build_where(year: int | None, search: str | None, status: str | None) -> tuple[str, list[Any]]:
    conditions: list[str] = []
    values: list[Any] = []
    if year is not None:
        conditions.append("tahun = ?")
        values.append(year)
    if status:
        conditions.append("status_lost = ?")
        values.append(status)
    if search and search.strip():
        pattern = f"%{search.strip()}%"
        conditions.append(
            "(perusahaan LIKE ? COLLATE NOCASE OR domain_lab LIKE ? COLLATE NOCASE "
            "OR bulan LIKE ? COLLATE NOCASE OR jenis_transaksi LIKE ? COLLATE NOCASE)"
        )
        values.extend([pattern, pattern, pattern, pattern])
    return (" WHERE " + " AND ".join(conditions)) if conditions else "", values


def _transaction_values(tx: TransactionInput) -> tuple[Any, ...]:
    values = tx.model_dump()
    values.update(
        derive_transaction_fields(
            tx.jenis_transaksi,
            tx.pengujian_stel,
            tx.laboratorium,
            tx.keterangan,
            tx.pph23,
        )
    )
    return tuple(values[column] for column in TRANSACTION_COLUMNS)


@app.get("/", include_in_schema=False)
def root(request: Request) -> Any:
    accept = request.headers.get("accept", "")
    if "application/json" in accept and "text/html" not in accept:
        return {"status": "Online", "message": "Telkom Revenue Database API Ready"}
    return RedirectResponse(url="https://telkom-intern-project.vercel.app", status_code=307)


@app.get("/api/health")
def health() -> dict[str, Any]:
    with database_session() as connection:
        count = connection.execute("SELECT COUNT(*) FROM transactions").fetchone()[0]
    return {"status": "ok", "transaction_count": count}


@app.post("/api/upload")
async def upload_excel(
    file: UploadFile = File(...),
    year: int | None = Query(default=None, ge=1900, le=2200),
) -> dict[str, Any]:
    filename = Path(file.filename or "upload.xlsx").name
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise HTTPException(status_code=415, detail="Gunakan workbook .xlsx atau .xlsm; format .xls lama tidak didukung.")

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(prefix="incoming_", suffix=suffix, dir=UPLOAD_DIR, delete=False) as temp_file:
            temp_path = Path(temp_file.name)
            size = 0
            while chunk := await file.read(1024 * 1024):
                size += len(chunk)
                if size > MAX_UPLOAD_BYTES:
                    raise HTTPException(status_code=413, detail="Ukuran file melebihi batas 50 MB.")
                temp_file.write(chunk)
        imported = process_excel_file(temp_path, year_override=year, source_name=filename)
    except HTTPException:
        raise
    except DuplicateWorkbookError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Workbook tidak dapat diproses: {exc}") from exc
    finally:
        await file.close()
        if temp_path and temp_path.exists():
            temp_path.unlink()

    return {
        "status": "Success" if imported else "Warning",
        "message": f"Berhasil menggabungkan {imported} transaksi dari '{filename}'." if imported else "Tidak ditemukan data transaksi yang valid.",
        "records_inserted": imported,
    }


@app.get("/api/analytics/summary")
def get_summary_metrics(
    year: int | None = Query(default=None, ge=1900, le=2200),
    q: str | None = Query(default=None, max_length=300),
    status: str | None = Query(default=None, max_length=100),
) -> dict[str, int | float]:
    where, values = _build_where(year, q, status)
    with database_session() as connection:
        row = connection.execute(
            f"SELECT COALESCE(SUM(nominal_revenue), 0), COALESCE(SUM(nominal_lost), 0), COUNT(*) FROM transactions{where}",
            values,
        ).fetchone()
    total_revenue, total_lost, total_transactions = row
    denominator = total_revenue + total_lost
    lost_ratio = total_lost / denominator * 100 if denominator else 0
    return {
        "total_revenue": float(total_revenue),
        "total_lost": float(total_lost),
        "total_transactions": int(total_transactions),
        "lost_ratio_percent": round(lost_ratio, 2),
    }


@app.get("/api/transactions")
def get_transactions(
    limit: int = Query(default=50, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
    year: int | None = Query(default=None, ge=1900, le=2200),
    q: str | None = Query(default=None, max_length=300),
    status: str | None = Query(default=None, max_length=100),
) -> dict[str, Any]:
    where, values = _build_where(year, q, status)
    with database_session() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM transactions{where}", values).fetchone()[0]
        rows = connection.execute(
            f"SELECT * FROM transactions{where} ORDER BY tahun DESC, id DESC LIMIT ? OFFSET ?",
            [*values, limit, offset],
        ).fetchall()
    return {"data": [dict(row) for row in rows], "total": total, "limit": limit, "offset": offset}


@app.get("/api/transactions/{tx_id}")
def get_transaction(tx_id: int) -> dict[str, Any]:
    with database_session() as connection:
        row = connection.execute("SELECT * FROM transactions WHERE id = ?", (tx_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail=f"Transaksi dengan id {tx_id} tidak ditemukan.")
    return dict(row)


@app.post("/api/transactions", status_code=201)
def create_transaction(tx: TransactionInput) -> dict[str, Any]:
    columns = ", ".join(f'"{column}"' for column in TRANSACTION_COLUMNS)
    placeholders = ", ".join("?" for _ in TRANSACTION_COLUMNS)
    with database_session() as connection:
        cursor = connection.execute(
            f"INSERT INTO transactions ({columns}) VALUES ({placeholders})",
            _transaction_values(tx),
        )
        transaction_id = cursor.lastrowid
        row = connection.execute("SELECT * FROM transactions WHERE id = ?", (transaction_id,)).fetchone()
    return dict(row)


@app.put("/api/transactions/{tx_id}")
def update_transaction(tx_id: int, tx: TransactionInput) -> dict[str, Any]:
    assignments = ", ".join(f'"{column}" = ?' for column in TRANSACTION_COLUMNS)
    with database_session() as connection:
        cursor = connection.execute(
            f"UPDATE transactions SET {assignments} WHERE id = ?",
            (*_transaction_values(tx), tx_id),
        )
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail=f"Transaksi dengan id {tx_id} tidak ditemukan.")
        row = connection.execute("SELECT * FROM transactions WHERE id = ?", (tx_id,)).fetchone()
    return {"status": "Updated", "data": dict(row)}


@app.delete("/api/transactions/{tx_id}")
def delete_transaction(tx_id: int) -> dict[str, Any]:
    with database_session() as connection:
        cursor = connection.execute("DELETE FROM transactions WHERE id = ?", (tx_id,))
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail=f"Transaksi dengan id {tx_id} tidak ditemukan.")
    return {"status": "Deleted", "id": tx_id}


@app.post("/api/chat")
def chat_with_data(query: ChatQuery) -> dict[str, str]:
    lower_prompt = query.prompt.strip().casefold()
    year_match = re.search(r"\b(20\d{2})\b", lower_prompt)
    target_year = int(year_match.group(1)) if year_match else None

    with database_session() as connection:
        if any(term in lower_prompt for term in ("belum bayar", "unpaid", "nunggak", "piutang")):
            year_clause = " AND tahun = ?" if target_year else ""
            parameters = [target_year] if target_year else []
            rows = connection.execute(
                "SELECT perusahaan, SUM(nominal_lost) AS total_lost FROM transactions WHERE nominal_lost > 0"
                + year_clause
                + " GROUP BY perusahaan ORDER BY total_lost DESC LIMIT 5",
                parameters,
            ).fetchall()
            if not rows:
                return {"reply": "Tidak ditemukan data piutang/unpaid untuk filter tersebut."}
            scope = str(target_year) if target_year else "Semua Tahun"
            lines = [f"Perusahaan dengan piutang/unpaid ({scope}):"]
            lines.extend(f"• {row['perusahaan']}: Rp {row['total_lost']:,.0f}" for row in rows)
            return {"reply": "\n".join(lines)}

        if any(term in lower_prompt for term in ("revenue", "pendapatan", "omzet")):
            year_clause = " WHERE tahun = ?" if target_year else ""
            parameters = [target_year] if target_year else []
            row = connection.execute(
                "SELECT COALESCE(SUM(nominal_revenue), 0), COUNT(*) FROM transactions" + year_clause,
                parameters,
            ).fetchone()
            scope = f" tahun {target_year}" if target_year else ""
            return {"reply": f"Total pendapatan{scope} adalah Rp {row[0]:,.0f} dari {row[1]:,} transaksi."}
        count = connection.execute("SELECT COUNT(*) FROM transactions").fetchone()[0]

    return {
        "reply": (
            f"Saya telah menganalisis {count:,} transaksi. Coba tanyakan total pendapatan "
            "(misalnya 'pendapatan 2024') atau daftar perusahaan yang belum bayar."
        )
    }


if __name__ == "__main__":
    import uvicorn

    app_import = "backend.main:app" if __package__ else "main:app"
    uvicorn.run(
        app_import,
        host=os.getenv("API_HOST", "127.0.0.1"),
        port=int(os.getenv("API_PORT", "8000")),
        reload=os.getenv("API_RELOAD", "true").lower() == "true",
    )
